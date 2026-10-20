"""
Resume text extraction and light structure detection.

Supports PDF, DOCX and plain text (the proposal calls for PDF/DOCX), from
either a file path (CLI) or raw bytes (API upload).
"""
from __future__ import annotations

import io
import re
from pathlib import Path

from docx import Document
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


class ResumeParseError(ValueError):
    """Raised when a resume cannot be read; the message is safe to show users."""


# --------------------------------------------------------------------------
# Text extraction
# --------------------------------------------------------------------------
def clean_text(text: str) -> str:
    """Normalise whitespace and bullet characters without changing content."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    text = re.sub(r"[\u2022\u25cf\u25aa\u25e6\u2023\u2043]", "•", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _pdf_text(stream: io.BytesIO) -> str:
    try:
        reader = PdfReader(stream)
        if reader.is_encrypted:
            raise ResumeParseError(
                "This PDF is password-protected. Remove the password and upload it again."
            )
        return "\n\n".join((page.extract_text() or "") for page in reader.pages)
    except ResumeParseError:
        raise
    except Exception as error:  # pypdf raises many different exception types
        raise ResumeParseError(
            "This PDF could not be read. Try exporting it again from your editor."
        ) from error


def _docx_text(stream: io.BytesIO) -> str:
    try:
        document = Document(stream)
    except Exception as error:
        raise ResumeParseError(
            "This DOCX file could not be read. Try saving it again as .docx."
        ) from error

    lines = [p.text for p in document.paragraphs]
    # Many resume templates put skills/dates in tables.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(dict.fromkeys(cells)))
    return "\n".join(lines)


def extract_text_from_bytes(data: bytes, filename: str) -> str:
    """Extract text from an uploaded resume. Raises ResumeParseError on problems."""
    suffix = Path(filename or "").suffix.lower()

    if suffix not in SUPPORTED_EXTENSIONS:
        raise ResumeParseError("Unsupported file type. Upload a PDF, DOCX, or TXT resume.")
    if not data:
        raise ResumeParseError("The uploaded file is empty.")

    stream = io.BytesIO(data)
    if suffix == ".pdf":
        raw = _pdf_text(stream)
    elif suffix == ".docx":
        raw = _docx_text(stream)
    else:
        raw = data.decode("utf-8", errors="replace")

    text = clean_text(raw)
    if len(text) < 40:
        raise ResumeParseError(
            "No readable text was found. If this is a scanned image, "
            "export a text-based PDF or upload a DOCX."
        )
    return text


def extract_resume_text(file_path: str) -> str:
    """Extract text from a resume on disk (used by the command-line tools)."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {file_path}")
    return extract_text_from_bytes(path.read_bytes(), path.name)


# --------------------------------------------------------------------------
# Section detection
# --------------------------------------------------------------------------
_SECTION_PATTERNS = {
    "skills": r"(technical\s+)?skills|technologies|tech stack|languages\s*&?\s*tools|core competencies",
    "experience": r"(professional\s+|work\s+|relevant\s+)?experience|employment|work history|internships?",
    "projects": r"(academic\s+|personal\s+|selected\s+)?projects?",
    "education": r"education|coursework|academics",
    "activities": r"activities|leadership|extracurriculars?|involvement|clubs|hackathons?|awards|honors",
    "summary": r"summary|objective|profile|about me",
}
_SECTION_RES = {
    name: re.compile(rf"^\s*(?:{pattern})\s*:?\s*$", re.IGNORECASE)
    for name, pattern in _SECTION_PATTERNS.items()
}


def detect_section(line: str) -> str | None:
    """Return the section name if `line` looks like a resume section heading."""
    stripped = line.strip()
    if not stripped or len(stripped) > 40:
        return None
    for name, regex in _SECTION_RES.items():
        if regex.match(stripped):
            return name
    return None


def split_sections(text: str) -> dict[str, list[str]]:
    """
    Group non-empty resume lines by section. Lines before the first heading go
    in "header". Unrecognised headings stay inside the previous section.
    """
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        heading = detect_section(line)
        if heading:
            current = heading
            sections.setdefault(current, [])
            continue
        sections[current].append(line)
    return {name: lines for name, lines in sections.items() if lines}
